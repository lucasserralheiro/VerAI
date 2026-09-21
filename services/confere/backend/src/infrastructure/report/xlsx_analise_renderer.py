"""T-516 a T-519 — O `Relatorio_Analise_Medição.xlsx` (ESPEC 009 §10).

Cinco abas: o quadro-resumo e as quatro situações. A forma reproduz o artefato de
referência — mesmos nomes de aba, mesma geometria de título, linha vazia e
cabeçalho — e acrescenta duas coisas que ele não tem e que um relatório em
planilha precisa ter: **as quantidades também como número**, para que a planilha
possa ser somada, e **as marcas**, para que as ressalvas viajem dentro do
arquivo.

Por que planilha e não um segundo documento: o `.docx` é o entregável formal e
continua sendo o único. A análise é papel de trabalho — quem a recebe filtra,
ordena e soma.
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from domain.entities.analysis import AnaliseDaMedicao, ItemClassificado, SituacaoDaAnalise
from domain.value_objects.classification import Classificacao
from domain.value_objects.quantity import NumberFormat, Quantity

# `R-XLS-01` — os nomes são os do artefato de referência. Renomeá-los quebraria a
# leitura de quem já recebe esse arquivo hoje, montado à mão.
ABA_DA_CLASSIFICACAO = {
    Classificacao.CRITICO: "Itens Críticos",
    Classificacao.MAIOR_RELEVANCIA: "Divergências Maior Relevância",
    Classificacao.DIVERGENTE: "Divergências",
    Classificacao.SEM_DIVERGENCIA: "Sem Divergência",
}

# Geometria do gabarito: título, linha vazia, cabeçalho, dados.
LINHA_DO_CABECALHO = 3
PRIMEIRA_LINHA_DE_DADO = 4

# As colunas são as do gabarito, na mesma ordem — ESPEC 013 `R-NUM-04`. Cada
# quantidade ocupa **uma** coluna, numérica, e a grafia do relatório vem do
# formato de exibição da célula, não de uma segunda coluna de texto.
COLUNAS = ("Código", "Descrição", "Contratado", "Medido", "Saldo")

# Largura por **nome de coluna**, não por letra — ESPEC 013, TASKS 013 §2.2.
#
# O mapa anterior era `{"A": 18, …, "I": 24}` e valia igual para as quatro abas
# de detalhe, que têm contagens de colunas diferentes: a marca caía em `I` numa
# aba e em `G` noutra, recebendo largura 24 num caso e 13 no outro — estreita
# para `Sim — perfil ou pacote`. Índice absoluto num arranjo que varia por aba
# erra em alguma delas por construção.
LARGURA_PADRAO = 13
LARGURAS = {"Código": 18, "Descrição": 62, "Contratado": 13, "Medido": 13, "Saldo": 13}
LARGURA_DA_MARCA = 26

TINTA_DO_CABECALHO = PatternFill("solid", fgColor="0E3D47")
FONTE_DO_CABECALHO = Font(bold=True, color="FFFFFF", size=10)
FONTE_DO_TITULO = Font(bold=True, size=12, color="0B2235")
FONTE_DA_GLOSA = Font(italic=True, size=9, color="4E747E")

# `R-NUM-03` e `R-NUM-07` — a grafia do relatório, por formato de célula.
#
# **Nada de `.##` aqui.** A primeira versão desta espec usou `#,##0.##` supondo
# que `#` fizesse sumir a casa decimal inteira quando não há dígito. O `#` faz
# sumir o **dígito**; o separador é literal e fica: `10` exibia `10,` na página
# aberta no Excel. Ver ESPEC 013 §2.5 e `R-NUM-08`.
#
# Os quatro códigos abaixo reproduzem `Quantity.formatar()` sem depender de
# nenhum truque do Excel — a decisão "tem casa decimal ou não" é tomada aqui,
# onde o valor está, e cada célula recebe o código que a sua grafia exige.
#
# **Escritos na convenção do OOXML**, que é inglesa: `.` decimal e `,` milhar. É
# o Excel quem os traduz para os separadores do idioma ao exibir. A grafia
# brasileira (`#.##0,00`) parece a certa aqui e produz arquivo inválido — quem
# vier "corrigir" vai quebrar o arquivo.
FORMATO_MILHAR_INTEIRO = "#,##0"
FORMATO_MILHAR_DECIMAL = "#,##0.00"
FORMATO_SIMPLES_INTEIRO = "0"
FORMATO_SIMPLES_DECIMAL = "0.00"

FORMATOS = (
    FORMATO_MILHAR_INTEIRO,
    FORMATO_MILHAR_DECIMAL,
    FORMATO_SIMPLES_INTEIRO,
    FORMATO_SIMPLES_DECIMAL,
)


def _formato_de(quantidade: Quantity) -> str:
    """`R-NUM-03` / `D-02` — o formato sai do domínio, não de uma tabela daqui.

    Duas perguntas, e as duas já respondidas fora deste arquivo:

    - **milhar ou simples** é atributo do item (`R-MED-04`), e o documento formal
      já o respeita. Uma segunda tabela aqui seria a segunda verdade para a mesma
      decisão, e as duas divergiriam no primeiro item novo;
    - **com casa decimal ou sem** é a mesma condição que `Quantity.formatar()`
      aplica — arredondar para duas casas e, se o resultado for redondo, não
      grafar decimal nenhum. Espelhá-la aqui é o que garante a paridade célula a
      célula, inclusive no `4,50` que a versão anterior declarava divergente.
    """
    quantizado = quantidade.valor.quantize(Decimal("0.01"))
    redondo = quantizado == quantizado.to_integral_value()

    if quantidade.formato is NumberFormat.MILHAR:
        return FORMATO_MILHAR_INTEIRO if redondo else FORMATO_MILHAR_DECIMAL
    return FORMATO_SIMPLES_INTEIRO if redondo else FORMATO_SIMPLES_DECIMAL

# `R-DOC-10` — determinismo. Sem fixar as datas, o pacote OOXML carrega o
# instante da geração em `docProps/core.xml` e dois arquivos idênticos em
# conteúdo diferem em bytes.
CARIMBO = datetime(2020, 1, 1, 0, 0, 0)


class XlsxAnaliseRenderer:
    """Implementa `IAnaliseRenderer` com `openpyxl`.

    `openpyxl` já é dependência do projeto — é ele que lê a planilha de medição.
    Nenhuma biblioteca nova entra por causa deste incremento.
    """

    def renderizar(self, analise: AnaliseDaMedicao, destino: Path) -> Path:
        livro = Workbook()
        livro.remove(livro.active)  # a aba que o openpyxl cria sozinho

        self._resumo(livro.create_sheet("Resumo Executivo"), analise)

        # `R-XLS-01` — as quatro saem **sempre**, inclusive vazias. Aba ausente
        # obrigaria quem recebe a distinguir "não havia itens" de "o relatório
        # não gerou a aba", e essa dúvida numa conferência é cara.
        for classificacao in Classificacao:
            self._detalhe(
                livro.create_sheet(ABA_DA_CLASSIFICACAO[classificacao]),
                analise.situacao(classificacao),
            )

        livro.properties.creator = "Confere"
        livro.properties.created = CARIMBO
        livro.properties.modified = CARIMBO

        destino.parent.mkdir(parents=True, exist_ok=True)
        livro.save(destino)
        return destino

    # ── Aba 1 · o quadro-resumo ───────────────────────────────────────────────

    def _resumo(self, planilha: Worksheet, analise: AnaliseDaMedicao) -> None:
        planilha["A1"] = "Resumo Executivo - Análise Medição"
        planilha["A1"].font = FONTE_DO_TITULO

        # `R-RES-03` — identificação do **dado da aplicação**. O gabarito grafa
        # `TC 52/SMIT/2024 - TA 02` e uma proposta que não existe em nenhuma das
        # duas fontes lidas: inventá-las aqui gravaria no código um dado que a
        # aplicação não tem como saber (PLANO 009 §6.2).
        planilha["A2"] = f"Contrato : {analise.contrato_referencia}"
        planilha["A3"] = f"Proposta : {analise.proposta_origem}"
        planilha["A4"] = f"Competência : {analise.competencia}"

        planilha["A6"] = "Classificação"
        planilha["B6"] = "Quantidade"
        planilha["C6"] = "Critério"
        for coluna in "ABC":
            celula = planilha[f"{coluna}6"]
            celula.font = FONTE_DO_CABECALHO
            celula.fill = TINTA_DO_CABECALHO

        linha = 7
        for situacao in analise.situacoes:
            planilha[f"A{linha}"] = situacao.rotulo
            planilha[f"B{linha}"] = situacao.quantidade
            planilha[f"C{linha}"] = situacao.glosa
            planilha[f"C{linha}"].font = FONTE_DA_GLOSA
            linha += 1

        # `R-RES-02` — a linha de total existe para que o invariante de
        # `R-ANA-05` fique visível a quem lê, e não só no teste.
        planilha[f"A{linha}"] = "Total de itens analisados"
        planilha[f"B{linha}"] = analise.total_itens
        planilha[f"A{linha}"].font = Font(bold=True)
        planilha[f"B{linha}"].font = Font(bold=True)

        planilha.column_dimensions["A"].width = 34
        planilha.column_dimensions["B"].width = 13
        planilha.column_dimensions["C"].width = 38

    # ── Abas 2 a 5 · as situações ─────────────────────────────────────────────

    def _detalhe(self, planilha: Worksheet, situacao: SituacaoDaAnalise) -> None:
        # `R-ANA-11` — *sem divergência* não traz saldo: ele é zero por definição
        # da própria situação, e uma coluna de zeros só ocupa largura.
        com_saldo = situacao.classificacao is not Classificacao.SEM_DIVERGENCIA
        marca = self._rotulo_da_marca(situacao.classificacao)

        planilha["A1"] = f"{situacao.rotulo} — {situacao.glosa}"
        planilha["A1"].font = FONTE_DO_TITULO

        colunas = [*COLUNAS] if com_saldo else [c for c in COLUNAS if c != "Saldo"]
        if marca:
            colunas.append(marca)

        for indice, titulo in enumerate(colunas, start=1):
            celula = planilha.cell(row=LINHA_DO_CABECALHO, column=indice, value=titulo)
            celula.font = FONTE_DO_CABECALHO
            celula.fill = TINTA_DO_CABECALHO
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for deslocamento, item in enumerate(situacao.itens):
            self._linha(planilha, PRIMEIRA_LINHA_DE_DADO + deslocamento, item, com_saldo, marca)

        # A largura acompanha o **nome** da coluna, onde quer que ele tenha caído.
        # Acrescentar ou remover uma coluna deixa de exigir tocar neste mapa.
        for indice, titulo in enumerate(colunas, start=1):
            largura = LARGURA_DA_MARCA if titulo == marca else LARGURAS.get(titulo, LARGURA_PADRAO)
            planilha.column_dimensions[get_column_letter(indice)].width = largura
        planilha.freeze_panes = planilha[f"A{PRIMEIRA_LINHA_DE_DADO}"]

    def _linha(
        self,
        planilha: Worksheet,
        linha: int,
        item: ItemClassificado,
        com_saldo: bool,
        marca: str,
    ) -> None:
        origem = item.linha
        planilha.cell(row=linha, column=1, value=str(origem.codigo))
        planilha.cell(row=linha, column=2, value=origem.descricao)

        # `R-NUM-01` — uma coluna por quantidade, numérica. A grafia do relatório
        # vem do formato de exibição (`R-NUM-03`): a célula guarda o valor cheio e
        # **mostra** o arredondado, que é o que uma planilha sabe fazer e o que
        # `R-XLS-05` tentava obter com uma segunda coluna de texto.
        quantidades = [origem.contratada, origem.medida]
        if com_saldo:
            quantidades.append(origem.saldo)

        for deslocamento, quantidade in enumerate(quantidades):
            # `Decimal` direto: `openpyxl` o aceita nativamente. Passar por
            # `float` traria erro binário para um número que instrui faturamento
            # — é a mesma recusa que `Quantity` faz no construtor.
            celula = planilha.cell(row=linha, column=3 + deslocamento, value=quantidade.valor)
            celula.number_format = _formato_de(quantidade)

        if marca:
            planilha.cell(
                row=linha,
                column=3 + len(quantidades),
                value=self._valor_da_marca(item, marca),
            )

    # ── As marcas que viajam dentro do arquivo ────────────────────────────────

    @staticmethod
    def _rotulo_da_marca(classificacao: Classificacao) -> str:
        """`R-XLS-04` — a marca de item sem previsão contratual, em coluna própria.

        O arquivo circula **sem a legenda da tela**, e sem esta coluna o item
        medido sem cobertura contratual seria só mais um crítico.

        **`Sem Divergência` não tem mais marca** (`R-XLS-03` revisada na v1.3):
        a coluna `Perfil ou pacote` era acréscimo nosso e o gabarito não a tem.
        A ressalva de §6.4 continua na tela e na resposta da API — que é onde
        ela nasceu —, e some do arquivo. Ver ESPEC 009 §14.
        """
        return "Sem previsão contratual" if classificacao is Classificacao.CRITICO else ""

    @staticmethod
    def _valor_da_marca(item: ItemClassificado, marca: str) -> str:
        """Texto legível, nunca `TRUE`.

        Quem abre a planilha lê a célula, não a documentação da coluna. `Sim` e
        vazio dizem o que é; `TRUE` e `FALSE` dizem que houve um booleano.
        """
        return "Sim — não consta no contrato" if item.sem_previsao_contratual else ""
