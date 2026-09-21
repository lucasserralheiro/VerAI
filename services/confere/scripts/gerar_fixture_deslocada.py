"""T-2034 — Fixture de códigos deslocados (ESPEC 027 §2.3).

Uma planilha no formato da aba `Levantamento`, mas com os códigos de serviço
na **coluna H** — fora das cinco primeiras colunas que ``LevantamentoReader``
lê (`COLUNAS_LIDAS = 5`). O resultado é uma leitura que não encontra nenhum
item, o cenário que a ESPEC 027 §1 mostra produzindo sessenta cartões na tela.

O cabeçalho não traz `Data do Levantamento :` nem `conforme contrato :` em
nenhuma das dez primeiras linhas — é o que faz `V-MED-02` disparar duas vezes,
parte dos sessenta.

Uso:
    python scripts/gerar_fixture_deslocada.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
DESTINO = RAIZ / "backend" / "tests" / "fixtures" / "levantamento_codigos_deslocados.xlsx"

COLUNA_DO_CODIGO = 8  # H — fora das cinco primeiras (COLUNAS_LIDAS = 5)

# Códigos plausíveis, só para preencher a coluna H. Não precisam casar com
# nenhum contrato real: `medicao.codigos` fica vazio de qualquer forma, porque
# o leitor nunca olha a coluna H.
CODIGOS = [
    "14.031.00020.00",
    "14.024.00006.00",
    "12.030.00001.00",
    "14.048.00027.00",
    "10.050.00001.00",
    "11.027.00003.00",
    "14.049.00038.00",
    "12.029.00021.00",
    "14.070.00002.00",
    "11.051.00012.00",
]


def main() -> None:
    livro = openpyxl.Workbook()
    aba = livro.active
    assert aba is not None
    aba.title = "Levantamento"

    aba.cell(row=1, column=1, value="LEVANTAMENTO - COMPROVAÇÃO FIXTURE DESLOCADA")
    # Linhas 2 a 4 propositalmente sem `Data do Levantamento :` nem
    # `conforme contrato :` — é o que faz a V-MED-02 disparar duas vezes.
    aba.cell(row=2, column=1, value="Planilha gerada para ESPEC 027 T-2034")

    for indice, codigo in enumerate(CODIGOS, start=1):
        linha = 4 + indice
        aba.cell(row=linha, column=1, value=f"Serviço de teste {indice}")
        aba.cell(row=linha, column=4, value="10")  # Quantidade Contratada
        aba.cell(row=linha, column=5, value="5")  # Quantidade Medida
        aba.cell(row=linha, column=COLUNA_DO_CODIGO, value=codigo)

    livro.save(DESTINO)
    livro.close()
    print(f"[ok] {DESTINO.name} — {len(CODIGOS)} códigos na coluna H, fora das colunas lidas")


if __name__ == "__main__":
    main()
