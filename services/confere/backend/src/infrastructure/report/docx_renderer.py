"""ESPEC 003 — Renderização do relatório em DOCX.

Implementa ``IReportRenderer``. O documento nasce de uma **cópia do modelo
institucional**, nunca montado do zero: é o que preserva as 7 fontes embutidas,
as 3 imagens, o cabeçalho e o rodapé (`R-DOC-01`).

Estrutura produzida:

    Seção 1 — retrato .... a capa do modelo, intocada (`R-DOC-02`)
    ╪ quebra de seção ╪
    Seção 2 — paisagem ... título, faixas e a tabela de 5 colunas
    ╪ quebra de seção ╪
    Seções 3 a 21 ....... um anexo de detalhamento cada, na orientação do
                          GRC (ESPEC 004 `R-ANX-02`, `R-ANX-03`)

Todas as seções herdam o cabeçalho e o rodapé por vínculo com a anterior, então
o papel timbrado aparece em todas as páginas sem duplicar recurso no pacote
(`R-DOC-05`).

**Os anexos entram no fim e não deslocam nada.** O teste-âncora continua medindo
as mesmas 55 linhas nas páginas 2 e 3 (`R-ANX-09`); se ele quebrar, o defeito
está aqui.
"""

from __future__ import annotations

import shutil
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

import docx
from docx.document import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from domain.entities.annex import Anexo, CelulaAnexo, ImagemAnexo, Orientacao
from domain.entities.report import Report, ReportLine
from infrastructure.report import layout, modelo, ooxml
from infrastructure.report.modelo import MODELO

# Unidade do OOXML para largura: um ponto tipográfico vale 20 twips.
TWIPS_POR_PONTO = 20.0

# Margens laterais da seção paisagem. Com 1,4 cm a largura útil dá exatamente a
# largura da tabela medida no relatório GRC — 26,9 cm. Manter as margens do
# modelo (1,9 cm) deixaria 1 cm a menos, e estreitar as colunas quebraria a
# fidelidade que é o critério de aceite (ESPEC 003 D-02).
MARGEM_LATERAL_PAISAGEM = Cm(1.4)

# Margens laterais das seções de anexo. 1,2 cm, e não 1,4: três anexos —
# `Servidores`, `ServicosVcloud` e `Comunicação Dados` — têm tabela mais larga
# que a área útil de 1,4 cm, por 0,6 a 1,4 mm. Reproduzir a largura do GRC exige
# esse milímetro; encolher a tabela para caber seria trocar a fidelidade que é o
# critério de aceite por uma margem redonda.
MARGEM_LATERAL_ANEXO = Cm(1.2)

# Cinco linhas no bloco de título: o título, uma em branco, as duas de rótulo e
# outra em branco antes da primeira seção. As duas vazias são medidas no
# relatório modelo — as faixas navy saltam de 57,1 para 72,0 e de 79,4 para
# 94,3, com a altura de linha em 7,4.
LINHAS_DO_TITULO = 5

# Carimbo de tempo fixo nas entradas do pacote. Um DOCX é um ZIP, e o ZIP grava
# a hora de gravação em cada entrada — duas execuções da mesma entrada sairiam
# com bytes diferentes só por isso, quebrando `R-DOC-10`. A data escolhida é
# arbitrária e sem significado: só precisa não variar.
CARIMBO_FIXO = (1980, 1, 1, 0, 0, 0)


def _normalizar_pacote(caminho: Path) -> None:
    """Reescreve o pacote com carimbo de tempo fixo, preservando a ordem.

    A ordem das entradas importa: ``[Content_Types].xml`` precisa vir primeiro
    para que o pacote seja reconhecido como OOXML.
    """
    with zipfile.ZipFile(caminho) as origem:
        entradas = [(info, origem.read(info.filename)) for info in origem.infolist()]

    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as destino:
        for info, dados in entradas:
            nova = zipfile.ZipInfo(info.filename, date_time=CARIMBO_FIXO)
            nova.compress_type = info.compress_type
            nova.external_attr = info.external_attr
            destino.writestr(nova, dados)


class DocxRenderer:
    """Renderiza o relatório sobre o papel timbrado da PRODAM."""

    def renderizar(self, relatorio: Report, destino: Path) -> Path:
        destino.parent.mkdir(parents=True, exist_ok=True)

        # Copiar antes de abrir deixa explícito que o modelo é entrada, nunca
        # saída: nada que aconteça aqui pode alterá-lo.
        shutil.copy2(MODELO, destino)
        documento = docx.Document(str(destino))

        self._preencher_a_capa(documento, relatorio)
        self._encerrar_a_capa(documento)
        self._abrir_secao_paisagem(documento)
        self._centralizar_timbrado(documento)
        self._preencher(documento, relatorio)

        documento.save(str(destino))
        _normalizar_pacote(destino)
        return destino

    # ── Estrutura ─────────────────────────────────────────────────────────────

    def _preencher_a_capa(self, documento: Document, relatorio: Report) -> None:
        """T-1410 / ESPEC 020 — a capa passa a nomear o contrato que ela comprova.

        Até aqui a capa era reproduzida como está no modelo (`R-DOC-02`), e o
        documento do PGM saía dizendo `SECRETARIA MUNICIPAL DE INOVAÇÃO E
        TECNOLOGIA` e `Contrato : TC 52/SMIT/2024 - TA 02`.

        **Escreve por posição, nunca por casamento de cadeia** (`R-CAP-03`,
        `D-09`): `PA-SMIT-260319-739` tem dois papéis no modelo — continuação da
        linha `Proposta :` e rodapé da capa —, e um `replace` escreveria a lista de
        propostas nos dois e a duplicaria numa delas.

        **Escreve nas duas cópias de cada caixa** (`D-02`). O `mc:AlternateContent`
        guarda uma representação moderna e uma de compatibilidade da mesma forma, e
        o Word escolhe conforme a versão: preencher uma só produz um documento que
        mostra o cliente certo numa máquina e o errado noutra.

        Só o **texto** muda. Fonte, corpo, cor, posição e a arte da capa ficam —
        é o que separa esta espec de reescrever a capa (`D-01`).
        """
        valores = {
            modelo.CAPA_CLIENTE: relatorio.cliente_da_capa,
            modelo.CAPA_SUBTITULO: relatorio.subtitulo_da_capa,
            modelo.CAPA_CONTRATO: f"Contrato : {relatorio.contrato_referencia}",
            # `R-CAP-07` — a lista inteira no primeiro nó, e o segundo esvazia: no
            # modelo ele era só a **quebra visual** de uma linha longa. Determinístico,
            # e sem matemática de layout.
            modelo.CAPA_PROPOSTAS: f"Proposta : {relatorio.propostas_da_capa}",
            modelo.CAPA_PROPOSTAS_CONTINUACAO: "",
            modelo.CAPA_PROPOSTAS_RODAPE: relatorio.propostas_da_capa,
        }

        for ordem, caixa in enumerate(documento.element.iter(qn("w:txbxContent"))):
            logica = ordem // modelo.CAPA_COPIAS_POR_CAIXA
            textos = list(caixa.iter(qn("w:t")))
            for indice, no in enumerate(textos):
                novo = valores.get((logica, indice))
                if novo is not None:
                    # `.text` do lxml escapa sozinho: um `&` no nome do órgão
                    # viraria `&amp;`. Escrever XML à mão aqui produziria pacote
                    # que o Word recusa inteiro.
                    no.text = novo

    def _encerrar_a_capa(self, documento: Document) -> None:
        """Remove a quebra de página e os parágrafos vazios que sobram do modelo.

        O modelo tem duas páginas: a capa e uma folha timbrada em branco, para
        quem for escrever nela à mão. Mantê-las produziria uma página vazia
        entre a capa e a tabela.

        A remoção começa na quebra de página e vai até o fim do corpo. Tudo o
        que vem depois dela no modelo é parágrafo vazio — a capa está antes, e
        não é tocada (`R-DOC-02`).
        """
        corpo = documento.element.body
        paragrafos = corpo.findall(qn("w:p"))

        quebra = f".//{qn('w:br')}[@{qn('w:type')}='page']"
        inicio = next((i for i, p in enumerate(paragrafos) if p.findall(quebra)), None)
        if inicio is None:
            return

        for paragrafo in paragrafos[inicio:]:
            corpo.remove(paragrafo)

    def _abrir_secao_paisagem(self, documento: Document) -> None:
        """T-206 — a seção do corpo, em paisagem."""
        capa = documento.sections[0]
        corpo = documento.add_section(WD_SECTION.NEW_PAGE)

        corpo.orientation = WD_ORIENT.LANDSCAPE
        corpo.page_width, corpo.page_height = capa.page_height, capa.page_width
        corpo.left_margin = corpo.right_margin = MARGEM_LATERAL_PAISAGEM
        corpo.top_margin = capa.top_margin
        corpo.bottom_margin = capa.bottom_margin

        # Cabeçalho e rodapé ficam vinculados à seção anterior — o padrão do
        # python-docx para seções novas. Declará-los aqui duplicaria recurso no
        # pacote sem ganho nenhum.

    def _centralizar_timbrado(self, documento: Document) -> None:
        """Centraliza as imagens do cabeçalho e do rodapé em relação à página.

        No modelo o cabeçalho é ancorado à **direita** da página e o rodapé à
        **esquerda**. Como as duas imagens têm cerca de 21 cm — a largura do A4
        retrato —, na capa elas ocupam a página inteira e o alinhamento não
        aparece. Na seção paisagem, de 29,7 cm, sobram 8,8 cm que se acumulam
        de um lado só, e o timbrado fica visivelmente torto.

        Centralizar resolve os dois casos com uma única mudança: no retrato a
        diferença é de milímetros, e no paisagem a sobra se divide.

        A alternativa seria esticar as imagens até a largura da página, mas isso
        levaria o cabeçalho de 4,5 para 6,4 cm de altura, invadindo o corpo
        (ESPEC 003 D-03). Deformar imagem institucional para resolver
        alinhamento é troca ruim.
        """
        for secao in documento.sections:
            for parte in (secao.header, secao.footer):
                for ancora in parte._element.findall(f".//{qn('wp:anchor')}"):
                    self._centralizar_ancora(ancora)

    @staticmethod
    def _centralizar_ancora(ancora: Any) -> None:
        """Troca a posição horizontal da âncora por `center`, relativa à página."""
        posicao = ancora.find(qn("wp:positionH"))
        if posicao is None:
            return

        posicao.set("relativeFrom", "page")
        for filho in list(posicao):
            posicao.remove(filho)

        alinhamento = OxmlElement("wp:align")
        alinhamento.text = "center"
        posicao.append(alinhamento)

    def _preencher(self, documento: Document, relatorio: Report) -> None:
        """T-209 a T-213 — o corpo: título, faixas, tabela e rodapé."""
        self._bloco_titulo(documento, relatorio)
        # `R-REL-05` — **sem agrupamento**. Era uma tabela por seção, com as
        # faixas navy de grupo e de seção e o cabeçalho de colunas repetido a
        # cada uma. Passa a ser uma tabela contínua, na ordem do contrato.
        if relatorio.linhas:
            self._bloco_de_linhas(documento, relatorio.linhas)
        # `D-06` — o que a aba conhece e o contrato não. Título neutro: o
        # documento vai ao órgão, e a ausência pode ser do PDF submetido, não do
        # contrato (ESPEC 018 §2.3, grupo C).
        #
        # ESPEC 028 `R-ZER-01` — menos as linhas que não afirmam quantidade
        # nenhuma. Calculado **uma vez** e passado adiante: o asterisco do título
        # e a nota do rodapé remetem a este bloco, e recomputar a condição em
        # cada ponto é como eles se separariam no dia em que a regra mudasse.
        bloco_final = self._bloco_final(relatorio)
        if bloco_final:
            self._bloco_de_linhas(documento, bloco_final, titulo=layout.TITULO_DEMAIS_ITENS)
        self._rodape(documento, relatorio, bloco_final)
        self._anexos(documento, relatorio)

    # ── Blocos ────────────────────────────────────────────────────────────────

    @staticmethod
    def _bloco_final(relatorio: Report) -> list[ReportLine]:
        """ESPEC 028 `R-ZER-01` / `R-ZER-02` — o bloco final **como ele sai**.

        A escolha de omitir é **daqui**, e não do caso de uso (`D-01`):
        `relatorio.demais_itens` continua inteiro para a tela, para a análise e
        para a API, que existem para conferir. Quem tem de encolher é a peça que
        vai ao órgão, onde uma linha de dois zeros só ocupa espaço.

        Vazio quando **todas** as linhas do bloco são zeradas: aí não há faixa,
        não há asterisco e não há nota (`R-ZER-04`) — o documento fica igual ao
        de quem nunca teve bloco final.
        """
        return [linha for linha in relatorio.demais_itens if not linha.sem_quantidade_alguma]

    def _bloco_titulo(self, documento: Document, relatorio: Report) -> None:
        """T-209 — faixa navy com título, data do levantamento e contrato.

        Três colunas, e não cinco: as três primeiras do relatório são fundidas
        à esquerda, e as duas de quantidade ficam à direita, como no modelo.
        """
        esquerda = sum(layout.LARGURAS_COLUNAS_DOCX[:3])
        larguras = (esquerda, layout.LARGURAS_COLUNAS_DOCX[3], layout.LARGURAS_COLUNAS_DOCX[4])

        data = f"{relatorio.data_levantamento:%d/%m/%Y}" if relatorio.data_levantamento else ""

        # As duas linhas do meio têm fundo navy **apenas nas duas colunas de
        # quantidade**; à esquerda o fundo é branco e o texto, preto. Medido no
        # relatório modelo: as faixas navy dessas linhas vão de 681 a 799, e não
        # de 37 a 799 como a do título.
        tabela = documento.add_table(rows=LINHAS_DO_TITULO, cols=3)
        ooxml.sem_bordas(tabela)
        ooxml.fixar_larguras(tabela, larguras)
        ooxml.altura_fixa(tabela, layout.ALTURA_LINHA)

        # Linha 0 — título, navy na largura toda.
        for celula in tabela.rows[0].cells:
            ooxml.sombrear(celula, layout.NAVY)
        ooxml.escrever(
            ooxml.mesclar_linha(tabela, 0), relatorio.titulo, negrito=True, cor=layout.BRANCO
        )

        # Linhas 2 e 3 — rótulos à esquerda em preto, cabeçalho das quantidades
        # em navy à direita.
        meio = (
            (f"Data do Levantamento : {data}", "Quantidade", "Quantidade"),
            (
                f"*Valores conforme contrato : {relatorio.contrato_referencia}",
                "Contratada*",
                "Medida",
            ),
        )
        for deslocamento, (esq, coluna4, coluna5) in enumerate(meio):
            primeira = deslocamento == 0
            fileira = tabela.rows[2 + deslocamento]

            ooxml.escrever(fileira.cells[0], esq, negrito=True)
            for celula, titulo in zip(fileira.cells[1:], (coluna4, coluna5), strict=True):
                ooxml.sombrear(celula, layout.NAVY)
                ooxml.escrever(
                    celula, titulo, negrito=True, cor=layout.BRANCO, alinhamento="center"
                )

            # O contorno envolve as duas linhas como um bloco só. O divisor
            # entre elas existe apenas à esquerda: no modelo esse traço vai até
            # x=681, deixando as duas colunas de quantidade sem separação.
            self._bordas_do_rotulo(fileira, primeira)

        # As linhas 1 e 4 ficam vazias e sem preenchimento: é assim no relatório
        # modelo, e é o respiro que separa o título das quantidades e as
        # quantidades da primeira seção.

    @staticmethod
    def _bordas_do_rotulo(fileira: Any, primeira: bool) -> None:
        """Contorno do bloco de rótulo, medido no relatório modelo.

        Os traços encontrados no GRC: horizontal em 71,8 e em 86,6 na largura
        toda — topo e base do bloco —, verticais nas duas extremidades, e um
        horizontal em 79,2 que **para em x=681**. É esse último que separa as
        duas linhas de rótulo sem cortar as colunas de quantidade.
        """
        esquerda, contratada, medida = fileira.cells

        if primeira:
            # Topo do bloco, laterais, e o divisor — só na porção à esquerda.
            ooxml.bordas_da_celula(esquerda, ("top", "left", "bottom"))
            ooxml.bordas_da_celula(contratada, ("top",))
            ooxml.bordas_da_celula(medida, ("top", "right"))
        else:
            # Base do bloco e laterais. Sem borda superior: o divisor da linha
            # de cima já a desenha, e à direita ele não existe de propósito.
            ooxml.bordas_da_celula(esquerda, ("left", "bottom"))
            ooxml.bordas_da_celula(contratada, ("bottom",))
            ooxml.bordas_da_celula(medida, ("bottom", "right"))

    def _bloco_de_linhas(
        self, documento: Document, linhas: list[ReportLine], titulo: str = ""
    ) -> None:
        """T-1232 / T-1233 — cabeçalho de colunas e as linhas, sem agrupamento.

        Uma tabela contínua. O cabeçalho de colunas é marcado como **linha de
        cabeçalho** (`R-REL-11` do OOXML: `tblHeader`), e por isso o Word o
        repete a cada quebra de página em vez de ele aparecer uma vez por seção.

        `titulo` é usado só pelo bloco final (`D-06`); no corpo principal não há
        faixa nenhuma.
        """
        faixa = 1 if titulo else 0
        tabela = documento.add_table(
            rows=faixa + 1 + len(linhas), cols=len(layout.LARGURAS_COLUNAS_DOCX)
        )
        ooxml.contornar(tabela, layout.GRADE)
        ooxml.fixar_larguras(tabela, layout.LARGURAS_COLUNAS_DOCX)
        ooxml.altura_fixa(tabela, layout.ALTURA_LINHA)

        if titulo:
            for celula in tabela.rows[0].cells:
                ooxml.sombrear(celula, layout.NAVY)
            ooxml.escrever(
                ooxml.mesclar_linha(tabela, 0), titulo, negrito=True, cor=layout.BRANCO
            )

        cabecalho = tabela.rows[faixa]
        for celula, rotulo in zip(cabecalho.cells, layout.CABECALHO_COLUNAS, strict=True):
            ooxml.escrever(celula, rotulo, negrito=True, alinhamento="center")
        ooxml.repetir_cabecalho(cabecalho)

        for deslocamento, linha in enumerate(linhas):
            celulas = tabela.rows[faixa + 1 + deslocamento].cells
            ooxml.escrever(celulas[0], str(linha.codigo))
            ooxml.escrever(celulas[1], linha.descricao)
            ooxml.escrever(celulas[2], linha.unidade)
            # T-1272 — célula vazia não é zero: onde a aba nada afirma sobre a
            # quantidade contratada, a célula sai em branco.
            contratada = linha.contratada.formatar() if linha.contratada_declarada else ""
            ooxml.escrever(celulas[3], contratada, alinhamento="right")
            # A coluna da medida é destacada em lavanda no modelo.
            ooxml.sombrear(celulas[4], layout.LAVANDA)
            ooxml.escrever(celulas[4], linha.medida.formatar(), alinhamento="right")

    # ── Anexos de detalhamento (ESPEC 004) ────────────────────────────────────

    def _anexos(self, documento: Document, relatorio: Report) -> None:
        """Uma seção por anexo, **depois** da tabela de comprovação (`R-ANX-01`).

        Nada antes disto muda: a capa e as páginas 2 e 3 já estão montadas, e o
        teste-âncora continua medindo as mesmas 55 linhas.
        """
        for anexo in relatorio.anexos:
            # ESPEC 036 `R-VAZ-01` — anexo sem conteúdo não vira página.
            #
            # **O teste antecede a criação da seção, e é aí que está a
            # correção** (`R-VAZ-03`). Até aqui `_secao_do_anexo` vinha antes:
            # apagar só o título e a frase deixaria a quebra de seção de pé e
            # trocaria a página que explica por uma **página em branco** — com
            # timbrado, cabeçalho e rodapé, que são vinculados à seção.
            #
            # O leitor continua trazendo os 19 anexos configurados, vazios
            # inclusive (`R-VAZ-04`): é o que mantém o fato disponível para a
            # `V-ANX-01`, que denuncia a planilha cujas abas têm outro nome.
            if anexo.vazio:
                continue
            secao = self._secao_do_anexo(documento, anexo)
            self._tabelas_do_anexo(documento, anexo, secao)

    def _secao_do_anexo(self, documento: Document, anexo: Anexo) -> Any:
        """Seção nova, na orientação que o GRC usa (`R-ANX-02`, `R-ANX-03`).

        A quebra de seção já leva o anexo para uma página nova — não é preciso
        acrescentar quebra de página. Cabeçalho e rodapé seguem vinculados à
        seção anterior, como na tabela de comprovação, e o timbrado aparece nas
        páginas do anexo sem duplicar recurso no pacote.
        """
        capa = documento.sections[0]
        secao = documento.add_section(WD_SECTION.NEW_PAGE)

        retrato = anexo.orientacao is Orientacao.RETRATO
        secao.orientation = WD_ORIENT.PORTRAIT if retrato else WD_ORIENT.LANDSCAPE
        if retrato:
            secao.page_width, secao.page_height = capa.page_width, capa.page_height
        else:
            secao.page_width, secao.page_height = capa.page_height, capa.page_width

        secao.left_margin = secao.right_margin = MARGEM_LATERAL_ANEXO
        secao.top_margin = capa.top_margin
        secao.bottom_margin = capa.bottom_margin
        return secao

    def _tabelas_do_anexo(self, documento: Document, anexo: Anexo, secao: Any) -> None:
        """Uma tabela, ou mais quando algum cabeçalho precisa se repetir.

        Os cortes vêm de `Anexo.cortes` (ESPEC 051 `R-SEG-02` a `R-SEG-06`),
        que aplica a restrição do Word: a marca de cabeçalho só vale nas
        primeiras fileiras de uma tabela. Cada corte abre um segmento novo —
        o preâmbulo (título, resumo, subtítulo da aba) fica no primeiro, e
        cada cabeçalho reconhecido abre o seu próprio, com sua própria marca.
        """
        larguras = self._larguras_do_anexo(anexo, secao)
        for inicio, fim, imagem in anexo.blocos():
            if imagem is not None:
                self._figura_do_anexo(documento, imagem, secao)
            else:
                self._faixa_de_tabelas(documento, anexo, larguras, inicio, fim)

    def _figura_do_anexo(self, documento: Document, imagem: ImagemAnexo, secao: Any) -> None:
        """A figura colada na aba, no tamanho com que o GRC a imprime.

        Sem limite de largura ela sairia com o dobro da página: o tamanho
        natural do PNG é o de uma captura de tela, não o de impressão.
        """
        util = secao.page_width.pt - secao.left_margin.pt - secao.right_margin.pt
        largura = imagem.largura_pt or util
        altura = imagem.altura_pt or 0.0
        if largura > util:
            altura = altura * util / largura
            largura = util

        paragrafo = documento.add_paragraph()
        paragrafo.paragraph_format.space_before = Pt(0)
        paragrafo.paragraph_format.space_after = Pt(0)
        paragrafo.add_run().add_picture(
            BytesIO(imagem.dados),
            width=Pt(largura),
            height=Pt(altura) if altura else None,
        )

    def _faixa_de_tabelas(
        self,
        documento: Document,
        anexo: Anexo,
        larguras: tuple[float, ...],
        inicio: int,
        fim: int,
    ) -> None:
        """Uma tabela por segmento — um a mais para cada cabeçalho reconhecido na faixa.

        ESPEC 051 `R-SEG-03`, `R-SEG-04` — generaliza o corte único da
        ESPEC 004: os cortes da faixa partem `[inicio, fim)` em segmentos, e
        cada segmento cujo início é uma linha de cabeçalho — primária ou
        adicional — recebe seu próprio `w:tblHeader`. Sem corte na faixa, um
        segmento só: o comportamento de sempre.

        `_tabela_do_anexo` não muda: ela já recebe `(inicio, fim)` por chamada
        e já calcula `forma_do_bloco` sobre esses dois argumentos — é aqui,
        passando intervalos menores, que a largura de cada segmento passa a
        ser a da sua própria tabela (`D-03`).
        """
        cortes = [c for c in anexo.cortes if inicio < c < fim]
        limites = [inicio, *cortes, fim]
        cabecalhos = {anexo.linha_cabecalho, *anexo.linhas_cabecalho_adicionais}

        for indice, segue_inicio in enumerate(limites[:-1]):
            segue_fim = limites[indice + 1]
            if indice:
                # Duas tabelas coladas: o Word junta as duas numa só quando não
                # há nada entre elas, e a junção desfaria o corte. Um
                # parágrafo de 1 pt separa sem abrir espaço perceptível.
                separador = documento.add_paragraph()
                separador.paragraph_format.space_before = Pt(0)
                separador.paragraph_format.space_after = Pt(0)
                corpo_separador = Pt(1)
                separador.add_run("").font.size = corpo_separador
                # ESPEC 053 `R-CEL-05` — a marca do parágrafo (`w:pPr/w:rPr`),
                # no mesmo corpo da execução acima. Sem ela, o "imperceptível"
                # herda os 12pt do documento e deixa de ser — é a marca, não a
                # execução, que o Word usa quando não há glifo para medir
                # (ESPEC 053 §2.5).
                marca = OxmlElement("w:rPr")
                tamanho_marca = OxmlElement("w:sz")
                tamanho_marca.set(qn("w:val"), str(int(corpo_separador.pt * 2)))
                marca.append(tamanho_marca)
                separador._p.get_or_add_pPr().append(marca)
            self._tabela_do_anexo(
                documento, anexo, larguras, segue_inicio, segue_fim,
                repetir=segue_inicio in cabecalhos,
            )

    def _tabela_do_anexo(
        self,
        documento: Document,
        anexo: Anexo,
        larguras: tuple[float, ...],
        inicio: int,
        fim: int,
        repetir: bool,
    ) -> None:
        linhas = anexo.linhas[inicio:fim]
        if not linhas:
            return

        # ESPEC 014 — a largura é do **bloco**, não da aba. O total de colunas
        # precisa estar decidido antes de `add_table`: a grade é resolvida uma
        # vez logo abaixo, e mexer nela depois é o caminho quadrático.
        forma = anexo.forma_do_bloco(inicio, fim)

        tabela = documento.add_table(rows=len(linhas), cols=forma.colunas)
        ooxml.contornar(tabela, layout.GRADE)
        # T-2013 / ESPEC 026 — a grade é resolvida **uma vez** e serve às duas:
        # `fixar_larguras` a percorria por conta própria, e o laço abaixo a
        # resolvia de novo. Resolver custa decidir `gridSpan` e `vMerge` em cada
        # célula da tabela, e `Usuários` tem 15 mil.
        grade = tabela._cells
        ooxml.fixar_larguras(tabela, larguras[: forma.colunas], grade)
        if anexo.altura_linha_pt:
            # Altura medida no GRC, como mínimo e não como valor exato. Altura
            # exata **corta** o que não couber, e o Word quebra linha onde o
            # Excel não quebrava: a linha cresce, o dado não some. Num documento
            # que instrui faturamento, essa é a troca certa.
            ooxml.altura_fixa(tabela, anexo.altura_linha_pt)

        # Tanto `Table.cell(i, j)` quanto `linha.cells` remontam a grade inteira
        # a cada acesso; usá-los dentro do laço tornaria `Usuários`, de 15 mil
        # células, quadrático — e na prática ele não terminava. A grade veio
        # resolvida de cima (T-2013).
        colunas = forma.colunas

        for indice, origem in enumerate(linhas):
            base = indice * colunas
            linha_vazia = (inicio + indice) in forma.linhas_vazias
            for coluna, celula in enumerate(origem[:colunas]):
                self._celula_do_anexo(
                    grade[base + coluna],
                    celula,
                    anexo.corpo,
                    sem_borda=linha_vazia or coluna in forma.colunas_vazias,
                )

        # As mesclagens vêm por último, e **por coordenada** (T-2008): a grade de
        # células acima já não é o que elas precisam. Aplicá-las em sequência é
        # seguro porque as regiões mescladas do Excel nunca se sobrepõem — fundir
        # uma não invalida a outra.
        self._mesclar_anexo(tabela, anexo, colunas, inicio, fim)

        if repetir:
            ooxml.repetir_cabecalho(tabela.rows[0])

    @staticmethod
    def _celula_do_anexo(
        alvo: Any, celula: CelulaAnexo, corpo: float, *, sem_borda: bool = False
    ) -> None:
        if celula.preenchimento:
            ooxml.sombrear(alvo, celula.preenchimento)
        if sem_borda or not celula.borda:
            # A grade é declarada na tabela inteira; onde a aba não emoldura,
            # ela é removida. É o que evita caixas vazias nas linhas em branco
            # que separam os blocos de cada aba.
            #
            # `sem_borda` acrescenta o caso em que a **aba emoldura e o GRC
            # não** (ESPEC 014 `R-BRD-03`, `R-BRD-04`): linha sem conteúdo e
            # coluna vazia que sobrou dentro do bloco.
            ooxml.sem_bordas_na_celula(alvo)
        # ESPEC 049 `R-CEL-01` — sempre, mesmo sem texto: célula vazia sem esta
        # normalização herda o `w:pPrDefault`/`w:rPrDefault` do modelo (8pt de
        # espaço depois do parágrafo, fonte 12pt), maior que o mínimo da linha
        # (`altura_linha_pt`) e maior que o corpo do próprio anexo — é o que
        # fazia o respiro entre faixas de título sair bem maior que a planilha.
        ooxml.escrever(
            alvo,
            celula.texto,
            negrito=celula.negrito,
            cor=celula.cor or layout.PRETO,
            corpo=corpo,
            alinhamento=celula.alinhamento,
        )

    @staticmethod
    def _mesclar_anexo(
        tabela: Any, anexo: Anexo, colunas: int, inicio: int, fim: int
    ) -> None:
        """T-2008 / ESPEC 026 — as coordenadas seguem, e não as duas células.

        Recebia `grade` e destruía o retângulo em duas células diagonais, para
        que `CT_Tc.merge` o redescobrisse por `_span_dimensions` — em O(fileiras)
        por mesclagem. `Office365` tem uma mesclagem por linha de dados, e no par
        do PGM isso respondia por 72 dos 82 s de perfil da renderização.

        As fileiras são resolvidas **uma vez por tabela**, o que é seguro porque
        mesclar remove `w:tc` e nunca `w:tr`. Quem guarda essa invariante é
        `test_desempenho.py::test_a_mesclagem_nao_varre_as_fileiras`.
        """
        fileiras = ooxml.fileiras_de(tabela)
        for mesclagem in anexo.mesclagens:
            if mesclagem.linha < inicio or mesclagem.ate_linha >= fim:
                continue
            if mesclagem.coluna >= colunas:
                continue
            # Limitada à última coluna do bloco, nunca descartada (ESPEC 014
            # `R-BRD-05`). Só uma mesclagem do piloto estoura — a faixa vazia de
            # `SOA` —, mas descartá-la deixaria a fileira em células soltas onde
            # a aba tem uma faixa só.
            ate_coluna = min(mesclagem.ate_coluna, colunas - 1)
            topo = mesclagem.linha - inicio
            ooxml.mesclar_regiao(
                fileiras,
                topo,
                mesclagem.coluna,
                altura=(mesclagem.ate_linha - inicio) - topo + 1,
                largura=ate_coluna - mesclagem.coluna + 1,
            )

    @staticmethod
    def _larguras_do_anexo(anexo: Anexo, secao: Any) -> tuple[float, ...]:
        """As larguras do GRC quando existem; a proporção da planilha quando não.

        Medir é melhor que proporcionalizar por um motivo concreto: as colunas
        da planilha são medidas em caracteres, e a razão entre elas não é a
        razão entre as larguras impressas. Dezessete dos dezenove anexos têm
        fronteira de coluna legível no PDF de referência e usam a medida direta.

        Nos outros dois a medição é ambígua — uma coluna sem borda não produz
        fronteira —, e a proporção da planilha é esticada até a **largura total
        medida**, que é inequívoca. Assim os dezenove saem com a largura de
        tabela do GRC, e dezessete com as colunas certas dentro dela.

        Não há folga a descontar: no OOXML a margem interna cabe **dentro** da
        largura da célula, então a soma das larguras é a largura da tabela. A
        folga da ESPEC 003 existia porque ali as medidas do GRC eram a área de
        texto, e não a caixa.
        """
        util = secao.page_width.pt - secao.left_margin.pt - secao.right_margin.pt

        medidas = anexo.larguras_medidas
        if not medidas:
            proporcoes = anexo.proporcoes or (1.0,) * anexo.total_colunas
            total = sum(proporcoes) or float(len(proporcoes))
            alvo = anexo.largura_total_pt or util
            medidas = tuple(alvo * proporcao / total for proporcao in proporcoes)

        # Rede de segurança: um anexo cuja tabela não coubesse na página sairia
        # com coluna cortada, que é o defeito que a ESPEC 003 já pagou uma vez.
        largura = sum(medidas)
        if largura > util:
            medidas = tuple(valor * util / largura for valor in medidas)
        return medidas

    def _rodape(
        self, documento: Document, relatorio: Report, bloco_final: list[ReportLine]
    ) -> None:
        """T-213 — contrato e proposta que originou as quantidades.

        Sem carimbo de relógio: uma data de geração quebraria a identidade entre
        duas execuções da mesma entrada, que é critério de aceite (`R-DOC-10`).

        `bloco_final` chega pronto de `_preencher`, e não é recalculado: a nota
        de `R-NOT-02` só existe porque a faixa com asterisco existe, e as duas
        têm de decidir pela mesma lista (ESPEC 028 `R-ZER-04`).
        """
        texto = f"Contrato: {relatorio.contrato_referencia}"
        if relatorio.proposta_origem:
            # ESPEC 019 `R-ADT-11` / `D-09` — todas as peças, com concordância.
            #
            # Saiu a palavra *Quantidades*: desde a ESPEC 018 `D-05` as duas
            # quantidades vêm da aba `Levantamento`, e o que estas propostas
            # originam é o **escopo** — quais itens existem, em que ordem e com
            # que designação. Dizer *quantidades* aqui apontava para a fonte errada.
            propostas = "as propostas" if len(relatorio.propostas) > 1 else "a proposta"
            texto += f"  ·  Conforme {propostas} {relatorio.proposta_origem}"

        paragrafo = documento.add_paragraph()
        execucao = paragrafo.add_run(texto)
        execucao.font.size = Pt(layout.CORPO_FONTE)
        execucao.font.name = layout.FONTE_REGULAR

        # ESPEC 024 `R-NOT-02` / `R-NOT-03` — só quando o bloco final existe.
        # Mesmo corpo e fonte do parágrafo acima (`R-NOT-04`), sem negrito e
        # sem cor nova.
        #
        # ESPEC 028 `R-ZER-04` — *existe* passa a querer dizer *tem linha
        # exibida*: `relatorio.demais_itens` cheio de linhas zeradas não rende
        # faixa nenhuma, e a nota apontaria para um asterisco que não foi escrito.
        if bloco_final:
            nota = documento.add_paragraph()
            execucao_nota = nota.add_run(layout.NOTA_DEMAIS_ITENS)
            execucao_nota.font.size = Pt(layout.CORPO_FONTE)
            execucao_nota.font.name = layout.FONTE_REGULAR
