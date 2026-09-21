"""ESPEC 036 — o anexo sem conteúdo não vira página.

Uma aba configurada e ausente da planilha produzia uma seção inteira dizendo
*"A planilha não trouxe conteúdo para este anexo."* — três páginas no par do PGM,
cinco no documento que originou a espec. `R-VAZ-01` as omite.

O critério é o da ESPEC 014 (`_tem_conteudo`: texto ou preenchimento, **borda
não conta**) aplicado ao anexo inteiro, mais a ausência de figura — `R-VAZ-02`.
Ele fecha, de passagem, dois defeitos que `not self.linhas` deixava passar: a aba
com resquício de formatação, que saía como página em branco, e a aba só com
figura, que saía com a frase e **sem a figura**.

O que a supressão apagaria junto — o sinal de planilha com o layout trocado —
volta pela `V-ANX-01`, e é por isso que ela está neste mesmo módulo.
"""

from __future__ import annotations

import base64
from datetime import date
from pathlib import Path
from typing import Any

import docx
import pytest

from domain.entities.annex import Anexo, CelulaAnexo, ImagemAnexo, Orientacao
from domain.entities.report import Report
from domain.entities.validation_finding import Severity, ValidationReport
from infrastructure.annex.anexo_reader import AnexoReader
from infrastructure.report.docx_renderer import DocxRenderer

SECOES_ANTES_DOS_ANEXOS = 2  # a capa e a tabela de comprovação
TABELAS_ANTES_DOS_ANEXOS = 1  # o bloco de título, que sai mesmo sem seções

# PNG de 1 × 1 transparente. O conteúdo não importa — o que se afirma é que a
# figura **chega ao pacote**; o tamanho impresso é assunto da ESPEC 004.
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _relatorio(anexos: list[Anexo]) -> Report:
    """Relatório sem seções: aqui o que está sob teste são os anexos."""
    return Report(
        titulo="LEVANTAMENTO - COMPROVAÇÃO",
        data_levantamento=date(2026, 7, 15),
        contrato_referencia="TC 52/SMIT/2024",
        proposta_origem="PA-SMIT-260319-739",
        anexos=anexos,
    )


def _documento(anexos: list[Anexo], destino: Path) -> Any:
    return docx.Document(str(DocxRenderer().renderizar(_relatorio(anexos), destino)))


def _textos(documento: Any) -> list[str]:
    return [paragrafo.text for paragrafo in documento.paragraphs]


# ── T-2289 · a aba com resquício de formatação (`R-VAZ-02` a) ─────────────────


def test_t2289_aba_com_linhas_e_nenhuma_celula_util_nao_vira_pagina(tmp_path: Path) -> None:
    """O caso que `not self.linhas` deixava passar, e que saía **em branco**.

    `AbaReader._colunas_uteis` apara as colunas finais sem conteúdo e pode
    devolver zero; as linhas sobrevivem, vazias. Medido contra o renderizador
    antes da correção: `vazio` dizia `False`, `total_colunas` era 0, e o anexo
    saía como uma tabela de 3 fileiras × 1 coluna, sem borda e sem texto — uma
    página em branco, sem sequer a frase.
    """
    anexo = Anexo(aba="SoResiduo", orientacao=Orientacao.RETRATO, corpo=8.0, linhas=((), (), ()))
    assert anexo.vazio

    documento = _documento([anexo], tmp_path / "residuo.docx")

    assert len(documento.sections) == SECOES_ANTES_DOS_ANEXOS
    assert len(documento.tables) == TABELAS_ANTES_DOS_ANEXOS


# ── T-2290 · a aba só com figura (`R-VAZ-02` b) ───────────────────────────────


def test_t2290_aba_so_com_figura_sai_com_a_figura(tmp_path: Path) -> None:
    """`blocos()` sempre soube emitir a figura; `vazio` curto-circuitava antes.

    Uma aba cujas células nada trazem e que tem figura ancorada saía com a frase
    e **sem o PNG** — o único caso em que a página de "não trouxe conteúdo"
    apagava conteúdo de verdade. Nenhuma das duas planilhas versionadas o
    exercita: `Internet` e `ServicosEmNuvem` têm figura *e* células.
    """
    anexo = Anexo(
        aba="SoFigura",
        orientacao=Orientacao.RETRATO,
        corpo=8.0,
        linhas=(),
        imagens=(ImagemAnexo(dados=PNG, linha=0, largura_pt=100.0, altura_pt=50.0),),
    )
    assert not anexo.vazio

    documento = _documento([anexo], tmp_path / "figura.docx")

    assert len(documento.inline_shapes) == 1
    assert len(documento.sections) == SECOES_ANTES_DOS_ANEXOS + 1


# ── T-2291 · borda não é conteúdo (`R-VAZ-02` c) ──────────────────────────────


def test_t2291_borda_nao_conta_como_conteudo(tmp_path: Path) -> None:
    """`R-BRD-02` continua valendo do mesmo lado.

    A ESPEC 014 mediu que a borda sobrevive em coluna que o Excel não imprime —
    as sete colunas sobrando de `Usuários` têm borda numa única linha. Usar a
    borda como sinal de conteúdo aqui manteria vivos justamente os anexos que
    esta espec existe para omitir.
    """
    so_borda = (CelulaAnexo(borda=True), CelulaAnexo(borda=True))
    anexo = Anexo(
        aba="SoBorda",
        orientacao=Orientacao.RETRATO,
        corpo=8.0,
        linhas=(so_borda, so_borda),
        proporcoes=(10.0, 10.0),
    )
    assert anexo.vazio

    documento = _documento([anexo], tmp_path / "borda.docx")

    assert len(documento.sections) == SECOES_ANTES_DOS_ANEXOS
    assert len(documento.tables) == TABELAS_ANTES_DOS_ANEXOS


# ── T-2291b · o oposto, que impede o critério de ser frouxo ───────────────────


def test_t2291_faixa_de_cor_sem_texto_e_conteudo(tmp_path: Path) -> None:
    """Preenchimento sem texto **é** conteúdo — é como as abas desenham cabeçalho.

    O par do teste acima. Um critério que exigisse texto omitiria a aba cujo
    cabeçalho é uma faixa `A52A2A` sem valor nas células, e essa é a forma de
    oito dos dezenove anexos.
    """
    faixa = (CelulaAnexo(preenchimento="A52A2A"),)
    anexo = Anexo(
        aba="SoFaixa",
        orientacao=Orientacao.RETRATO,
        corpo=8.0,
        linhas=(faixa,),
        proporcoes=(10.0,),
    )
    assert not anexo.vazio

    documento = _documento([anexo], tmp_path / "faixa.docx")

    assert len(documento.sections) == SECOES_ANTES_DOS_ANEXOS + 1


# ── T-2292 · `V-ANX-01` — um achado, não dezenove ─────────────────────────────


def _vazio(aba: str) -> Anexo:
    return Anexo(aba=aba, orientacao=Orientacao.RETRATO, corpo=8.0)


def _com_conteudo(aba: str) -> Anexo:
    return Anexo(
        aba=aba,
        orientacao=Orientacao.RETRATO,
        corpo=8.0,
        linhas=((CelulaAnexo(texto="x"),),),
        proporcoes=(10.0,),
    )


def test_t2292_nenhuma_aba_reconhecida_produz_um_achado() -> None:
    """`R-GRD-06` — a causa, uma vez, e não a consequência dezenove vezes.

    Dezenove anexos vazios não são dezenove problemas: são um só, e é a planilha
    nomear as abas de outro jeito. Foi assim que nasceram os 57 achados do
    `PA-PGM` que a ESPEC 017 teve de desfazer.
    """
    from infrastructure.validations.annex_validations import (
        v_anx_01_nenhuma_aba_de_anexo_reconhecida,
    )

    achados = ValidationReport()
    v_anx_01_nenhuma_aba_de_anexo_reconhecida([_vazio(f"Aba{i}") for i in range(19)], achados)

    assert len(achados.achados) == 1
    assert achados.achados[0].validacao == "V-ANX-01"
    assert achados.achados[0].severidade is Severity.AVISA
    assert not achados.bloqueado


def test_t2292_ausencia_parcial_nao_avisa() -> None:
    """`D-04` — 3 de 19 no PGM é o escopo contratado, não anomalia.

    Aquele órgão não contratou colocation. Avisar a cada competência
    transformaria um fato estável em ruído mensal, e é este teste que impede a
    regra de escorregar para lá numa refatoração.
    """
    from infrastructure.validations.annex_validations import (
        v_anx_01_nenhuma_aba_de_anexo_reconhecida,
    )

    achados = ValidationReport()
    anexos = [_vazio(f"Aba{i}") for i in range(18)] + [_com_conteudo("Detalhes")]
    v_anx_01_nenhuma_aba_de_anexo_reconhecida(anexos, achados)

    assert achados.achados == []


def test_t2292_sem_anexos_lidos_nao_ha_o_que_afirmar() -> None:
    """Relatório bloqueado não lê anexo nenhum — e ausência de leitura não é achado.

    `R-GRD-06`: a validação continua chamável fora do fluxo do contêiner, e fora
    dele a lista vazia significa *"não se leu"*, não *"não veio nada"*.
    """
    from infrastructure.validations.annex_validations import (
        v_anx_01_nenhuma_aba_de_anexo_reconhecida,
    )

    achados = ValidationReport()
    v_anx_01_nenhuma_aba_de_anexo_reconhecida([], achados)

    assert achados.achados == []


# ── T-2293 · o par do PGM, que é onde o caso é real ───────────────────────────


@pytest.fixture(scope="module")
def anexos_do_pgm(caminho_levantamento_pgm: Path) -> list[Anexo]:
    return AnexoReader().ler(caminho_levantamento_pgm)


def test_t2293_as_tres_abas_ausentes_do_pgm_sao_as_medidas(anexos_do_pgm: list[Anexo]) -> None:
    """A régua da entrega, do lado do leitor (`R-VAZ-04`).

    O leitor continua trazendo os **19** anexos configurados, vazios inclusive —
    é o que mantém o fato disponível para a `V-ANX-01`. Quem decide a página é o
    renderizador.
    """
    assert len(anexos_do_pgm) == 19
    assert [anexo.aba for anexo in anexos_do_pgm if anexo.vazio] == [
        "Colocation",
        "Comunicação Dados",
        "CertificadosDigitais",
    ]


def test_t2293_o_documento_do_pgm_perde_tres_secoes_e_nenhuma_tabela(
    documento_do_pgm: Path, anexos_do_pgm: list[Anexo]
) -> None:
    """**34 tabelas antes da ESPEC 051, 36 depois** — e é essa a asserção que importa.

    Seção é página; tabela é conteúdo. As três abas ausentes não produziam
    tabela nenhuma, então esse número não pode mudar — nem para menos, que
    seria anexo perdido, nem para mais.

    Medido no par completo do PGM, com aditivo, em 2026-08-31 (ESPEC 036):
    **21 → 18 seções**, **3 → 0** frases, **34 tabelas nos dois estados**. O
    "antes" saiu da prova por desligamento da `T-2302`, e não de estimativa —
    a primeira medição contou 32 porque foi feita sobre uma renderização **só
    dos anexos**, sem o bloco de título nem a tabela de comprovação.

    **34 → 36 em 2026-09-10 (ESPEC 051), medido e não estimado.** `Servidores`
    e `ServidoresSemDesenv` — abas com uma segunda tabela, de forma diferente,
    empilhada — ganharam uma tabela a mais cada, uma por anexo: são as duas
    únicas do catálogo com `cabecalhos_adicionais`. Seções e frases **não se
    movem** — nenhum anexo aparece ou desaparece, só um deles se desenha com
    mais uma tabela.

    Usa a fixture de sessão do `conftest`: renderizar o PGM de novo custaria o
    que a `T-2001` existe para não pagar duas vezes.
    """
    documento = docx.Document(str(documento_do_pgm))
    com_conteudo = [anexo for anexo in anexos_do_pgm if not anexo.vazio]

    assert len(com_conteudo) == 16
    assert len(documento.sections) == SECOES_ANTES_DOS_ANEXOS + 16
    assert not [texto for texto in _textos(documento) if "não trouxe conteúdo" in texto]
    assert len(documento.tables) == 36


def test_t2293_os_dezesseis_saem_na_ordem_de_anexos_json(
    documento_do_pgm: Path, anexos_do_pgm: list[Anexo]
) -> None:
    """`R-VAZ-06` — sem renumeração e sem marca de que houve omissão.

    A orientação de cada seção é o que o documento guarda da ordem: catorze
    retratos e cinco paisagens no piloto, e aqui a sequência dos 16 que sobram.
    """
    documento = docx.Document(str(documento_do_pgm))
    secoes = list(documento.sections)[SECOES_ANTES_DOS_ANEXOS:]

    saiu = [
        Orientacao.PAISAGEM if secao.page_width > secao.page_height else Orientacao.RETRATO
        for secao in secoes
    ]
    assert saiu == [anexo.orientacao for anexo in anexos_do_pgm if not anexo.vazio]
