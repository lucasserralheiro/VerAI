"""ESPEC 028 — a linha do bloco final que não diz nada não vira linha.

Os dois pares reais provam a regra **acontecendo** (`test_anchor_por_codigo`):
o piloto perde três das quatro linhas do bloco final, o PGM perde sete das treze.
O que eles não trazem é o caso extremo — bloco final **inteiro** zerado — nem o
da célula em branco, que a `R-ZER-03` separa do zero de propósito e que nenhuma
das duas planilhas exercita no bloco final.

São cenários construídos, e é por isso que eles moram aqui e não junto do âncora:
o âncora afirma sobre documento medido, este módulo afirma sobre regra.

Cada relatório é renderizado **sem anexos** — o custo da geração é escrever os
anexos (conftest §"O que é caro"), e nenhuma asserção daqui olha para eles.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import docx
import pytest
from leitura_relatorio import ler

from domain.entities.report import Report, ReportLine
from domain.value_objects.quantity import Quantity
from domain.value_objects.service_code import ServiceCode
from infrastructure.report import layout
from infrastructure.report.docx_renderer import DocxRenderer


def _linha(codigo: str, contratada: str, medida: str, *, declarada: bool = True) -> ReportLine:
    return ReportLine(
        codigo=ServiceCode(codigo),
        descricao=f"ITEM {codigo}",
        unidade="UNIDADE",
        contratada=Quantity(Decimal(contratada)),
        medida=Quantity(Decimal(medida)),
        contratada_declarada=declarada,
    )


def _relatorio(
    linhas: list[ReportLine] | None = None, demais: list[ReportLine] | None = None
) -> Report:
    return Report(
        titulo="LEVANTAMENTO - COMPROVAÇÃO SMIT SUSTENTAÇÃO - CATÁLOGO DE SERVIÇOS DIT",
        data_levantamento=date(2026, 7, 15),
        contrato_referencia="TC 52/SMIT/2024",
        proposta_origem="PA-SMIT-260319-739",
        propostas=("PA-SMIT-260319-739",),
        cliente="SECRETARIA MUNICIPAL DE INOVAÇÃO E TECNOLOGIA",
        linhas=list(linhas or []),
        demais_itens=list(demais or []),
    )


def _renderizar(relatorio: Report, destino: Path) -> Path:
    return DocxRenderer().renderizar(relatorio, destino / "relatorio.docx")


# ── `R-ZER-01` e `R-ZER-03` · a afirmação sobre a linha ───────────────────────


@pytest.mark.parametrize(
    ("contratada", "medida", "declarada", "esperado"),
    [
        ("0", "0", True, True),
        # Um dos dois zeros basta para a linha continuar existindo: é
        # literalmente o pedido de origem desta espec.
        ("0", "2", True, False),
        ("5", "0", True, False),
        ("1", "1", True, False),
        # `R-ZER-03` — célula vazia não é zero (T-1272). A aba nada afirma sobre
        # o contratado, e um silêncio não é o segundo zero.
        ("0", "0", False, False),
    ],
)
def test_sem_quantidade_alguma_exige_dois_zeros_afirmados(
    contratada: str, medida: str, declarada: bool, esperado: bool
) -> None:
    linha = _linha("14.049.00004.00", contratada, medida, declarada=declarada)
    assert linha.sem_quantidade_alguma is esperado


# ── `R-ZER-01` · no bloco final, a linha zerada não é desenhada ───────────────


def test_a_linha_zerada_do_bloco_final_nao_vai_ao_documento(tmp_path: Path) -> None:
    relatorio = _relatorio(
        linhas=[_linha("11.027.00001.00", "10", "10")],
        demais=[
            _linha("12.029.00001.00", "0", "0"),
            _linha("14.049.00054.00", "0", "2"),
            _linha("14.024.00001.00", "0", "0"),
        ],
    )
    gerado = _renderizar(relatorio, tmp_path)

    codigos = [linha.codigo for linha in ler(gerado)]
    assert codigos == ["11.027.00001.00", "14.049.00054.00"]


def test_a_linha_com_um_zero_so_continua_no_documento(tmp_path: Path) -> None:
    """`R-ZER-01` pela contrapositiva — as duas metades do pedido de origem.

    `0 / 2` é consumo sem cobertura contratual e `5 / 0` é contratado que não se
    consumiu: nenhum dos dois é ausência de assunto, e os dois seguem no
    documento.
    """
    relatorio = _relatorio(
        demais=[_linha("14.049.00054.00", "0", "2"), _linha("14.071.00006.00", "5", "0")]
    )
    gerado = _renderizar(relatorio, tmp_path)

    assert {linha.codigo for linha in ler(gerado)} == {
        "14.049.00054.00",
        "14.071.00006.00",
    }


# ── `R-ZER-02` · o bloco do contrato não é tocado ─────────────────────────────


def test_a_linha_zerada_do_contrato_permanece(tmp_path: Path) -> None:
    """`R-ZER-02` — ali o zero é afirmação do contrato, e prova cobertura.

    O item que o contrato prevê e o período não consumiu é resposta a *"e o item
    X?"*. Some-lo obrigaria quem confere a voltar ao contrato para descobrir que
    a resposta era zero — que é o oposto do que o documento existe para fazer.
    """
    relatorio = _relatorio(linhas=[_linha("11.027.00001.00", "0", "0")])
    gerado = _renderizar(relatorio, tmp_path)

    assert [linha.codigo for linha in ler(gerado)] == ["11.027.00001.00"]


# ── `R-ZER-04` · bloco final inteiro zerado é bloco final nenhum ──────────────


def test_bloco_final_todo_zerado_nao_deixa_faixa_asterisco_nem_nota(tmp_path: Path) -> None:
    """O documento fica idêntico ao de quem nunca teve bloco final.

    O risco que este teste guarda é o do asterisco órfão: a `R-NOT-01` e a
    `R-NOT-02` da ESPEC 024 condicionavam-se a `relatorio.demais_itens`, que
    aqui **não é vazio** — a faixa é que não chega a existir.
    """
    relatorio = _relatorio(
        linhas=[_linha("11.027.00001.00", "10", "10")],
        demais=[_linha("12.029.00001.00", "0", "0"), _linha("14.024.00001.00", "0", "0")],
    )
    gerado = _renderizar(relatorio, tmp_path)
    documento = docx.Document(str(gerado))

    faixas = [tabela.rows[0].cells[0].text.strip() for tabela in documento.tables]
    assert layout.TITULO_DEMAIS_ITENS not in faixas
    assert layout.NOTA_DEMAIS_ITENS not in [p.text for p in documento.paragraphs]

    # E o corpo continua sendo o do contrato: a regra tira as linhas do bloco
    # final, não o documento.
    assert [linha.codigo for linha in ler(gerado)] == ["11.027.00001.00"]


def test_com_uma_linha_sobrevivente_a_faixa_e_a_nota_voltam(tmp_path: Path) -> None:
    """O par da asserção anterior — `R-ZER-04` nos dois sentidos."""
    relatorio = _relatorio(
        demais=[_linha("12.029.00001.00", "0", "0"), _linha("14.049.00054.00", "0", "2")]
    )
    gerado = _renderizar(relatorio, tmp_path)
    documento = docx.Document(str(gerado))

    faixas = [tabela.rows[0].cells[0].text.strip() for tabela in documento.tables]
    assert layout.TITULO_DEMAIS_ITENS in faixas
    assert layout.NOTA_DEMAIS_ITENS in [p.text for p in documento.paragraphs]