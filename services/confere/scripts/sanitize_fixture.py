"""T-05 e T-302 — Gera a fixture sanitizada da planilha de medição.

A planilha original traz ``Usuários`` com 1.021 registros nominais de servidores
públicos e ``Office365`` com 363 — login, nome completo, e-mail institucional.
Fixá-los no histórico do Git é impraticável de desfazer, e por isso eles nunca
entram (PLANO 001 D-04).

**O que mudou com os anexos (ESPEC 004 `R-ANX-10`):** essas duas abas eram
simplesmente removidas, porque o MVP só lia ``Levantamento``. Agora elas são
dois dos dezenove anexos — e logo os dois mais difíceis, por volume e por
paginação. Removê-las deixaria os testes sem cobrir justamente o que pode
quebrar.

Então elas **ficam, com dados sintéticos**: mesma estrutura, mesma contagem de
linhas e colunas, mesma formatação e mesmas mesclagens; só as colunas
identificadoras são substituídas. Os testes verificam forma, contagem e
paginação, que é o que precisam verificar.

A substituição é determinística: rodar duas vezes produz a mesma fixture.

Uso:
    python scripts/sanitize_fixture.py
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.cell_range import CellRange

RAIZ = Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / "docs" / "documentos"
DESTINO = RAIZ / "backend" / "tests" / "fixtures"

PLANILHA = "SMIT SUSTENTAÇÃO_Levantamento_05969_TC 52SMIT2024_15072026_101414_V2.0.xlsx"
CONTRATO = "PA-SMIT-260319-739 Q-00739-7.pdf"
MODELO = "SMIT_SUSTENTACAO_Levantamento_05969_TC_52SMIT2024_15072026_101414_V2.0___GRC.pdf"

# T-1290 — o segundo par, do PGM (ESPEC 018). A planilha traz `Usuários` com
# **997** registros nominais, e por isso segue o mesmo caminho da do SMIT.
PLANILHA_PGM = "PMG/PGM_TC 015_Levantamento_06008_TC 015PGM2024_23072026_095059_V1.0.xlsx"
CONTRATO_PGM = "PMG/PA-PGM-251015-159 v5.0.pdf"

# T-1305 — o aditivo do PGM (ESPEC 019). Três tabelas de itens em duas páginas,
# com três geometrias distintas, e **nenhuma página com exatamente oito
# divisórias** — é o documento que a `R-GRD-02` anterior não lia (§2.3).
ADITIVO_PGM = "PMG/PA-PGM-260304-715 - Q-00715-5_aditivo.pdf"

# Abas com dado pessoal identificável. **A linha do cabeçalho é derivada**, e
# não declarada: é a primeira linha da aba cujos rótulos incluem alguma coluna
# identificadora.
#
# Era um par `aba: linha` medido no arquivo do SMIT — `Office365: 17`. No
# arquivo do PGM o mesmo cabeçalho está na linha **22**, e a constante do SMIT
# teria deixado 1.488 linhas de nomes e e-mails passarem intactas. Uma
# coordenada medida num documento é uma afirmação sobre aquele documento
# disfarçada de afirmação sobre o formato — a mesma lição da ESPEC 017 `D-03`.
ABAS_COM_DADO_PESSOAL = ("Usuários", "Office365")

# A aba que o MVP consome, e a única onde a mesclagem das faixas precisa ser
# preservada na fixture (ver `_materializar_faixas`).
ABA_DE_ITENS = "Levantamento"

# Até onde procurar o cabeçalho. Acima disso é bloco de resumo, não lista.
LINHAS_DE_BUSCA_DO_CABECALHO = 30

# Rótulos de coluna cujo conteúdo identifica uma pessoa. A substituição é feita
# por **rótulo**, não por posição: uma coluna a mais na competência seguinte
# deslocaria os índices e deixaria nome real passar.
COLUNAS_IDENTIFICADORAS = {"login", "nome", "email", "e-mail", "usuário", "usuario"}

# Domínio reservado pela RFC 2606 para exemplos — nunca resolve, e deixa óbvio
# na leitura do arquivo que o endereço é fabricado.
DOMINIO_SINTETICO = "exemplo.invalid"

_PRENOMES = (
    "ANA", "BRUNO", "CARLA", "DANIEL", "ELISA", "FABIO",
    "GISELE", "HELIO", "IARA", "JOAO", "LUCIA", "MARCOS",
)
_SOBRENOMES = (
    "ALMEIDA", "BARBOSA", "CARDOSO", "DIAS", "ESTEVES",
    "FONSECA", "GOMES", "HENRIQUES", "IZIDORO", "JARDIM",
)


def _nome_sintetico(sequencia: int) -> str:
    prenome = _PRENOMES[sequencia % len(_PRENOMES)]
    sobrenome = _SOBRENOMES[(sequencia // len(_PRENOMES)) % len(_SOBRENOMES)]
    return f"{prenome} {sobrenome}"


def _login_sintetico(sequencia: int) -> str:
    return f"u{sequencia:06d}"


def _linhas_repetidas_na_origem(origem: Path) -> set[int]:
    """As linhas que o leitor **recebe** com o mesmo texto em mais de uma coluna.

    Lidas com a mesma abertura da produção — ``read_only=True`` —, porque é
    exatamente aí que o comportamento diverge entre os dois arquivos: o do PGM
    entrega a faixa mesclada repetida nas cinco colunas, e o do SMIT entrega em
    coluna única.

    Materializar **todas** as mesclagens tornaria a fixture do piloto diferente
    do arquivo do piloto — 85 faixas passariam de 1/5 para 5/5, e o portão que
    afirma *"o piloto não mudou"* deixaria de significar isso.
    """
    livro = openpyxl.load_workbook(origem, read_only=True, data_only=True)
    try:
        if ABA_DE_ITENS not in livro.sheetnames:
            return set()
        repetidas = set()
        for numero, linha in enumerate(livro[ABA_DE_ITENS].iter_rows(values_only=True), start=1):
            preenchidas = [str(v).strip() for v in linha if v is not None and str(v).strip()]
            if len(preenchidas) > 1 and len(set(preenchidas)) == 1:
                repetidas.add(numero)
        return repetidas
    finally:
        livro.close()


def _materializar_faixas(livro: Any, linhas_repetidas: set[int]) -> int:
    """T-1290 — preserva na fixture a forma com que o leitor recebe as faixas.

    A faixa de bloco do levantamento do PGM é célula mesclada, e o modo
    somente-leitura do ``openpyxl`` a entrega com o texto **repetido nas cinco
    colunas**. Foi essa forma que o crivo de faixa rejeitava, desligando a
    ``R-MED-02`` no PGM inteiro (ESPEC 018 §2.8).

    Ao regravar o arquivo, o ``openpyxl`` faz o contrário: guarda a mesclagem e
    **apaga** as células não-âncora. Medido — a faixa sai de 5/5 preenchidas no
    original para 1/5 na fixture.

    O efeito é o pior possível numa fixture: ela **normaliza o insumo e some com
    o defeito**. Uma suíte inteira passaria sobre um arquivo que já não tem a
    forma do arquivo real — que é, letra por letra, a história deste projeto.

    Aqui a mesclagem é desfeita e o valor materializado em todas as células do
    intervalo — **e apenas nas linhas que a origem já entregava repetidas**
    (``_linhas_repetidas_na_origem``). A mesclagem é cosmética; a forma que o
    leitor recebe, não.
    """
    if ABA_DE_ITENS not in livro.sheetnames or not linhas_repetidas:
        return 0

    planilha = livro[ABA_DE_ITENS]
    materializadas = 0
    for intervalo in list(planilha.merged_cells.ranges):
        limites = CellRange(str(intervalo))
        if limites.min_row not in linhas_repetidas:
            continue
        ancora = planilha.cell(limites.min_row, limites.min_col).value
        if ancora is None:
            continue
        planilha.unmerge_cells(str(intervalo))
        for linha in range(limites.min_row, limites.max_row + 1):
            for coluna in range(limites.min_col, limites.max_col + 1):
                planilha.cell(linha, coluna, ancora)
        materializadas += 1
    return materializadas


def _linha_do_cabecalho(planilha: Any) -> int | None:
    """A primeira linha cujos rótulos incluem alguma coluna identificadora.

    Derivada do próprio arquivo (ver ``ABAS_COM_DADO_PESSOAL``). Devolve
    ``None`` quando a aba não tem coluna identificadora nenhuma — caso em que
    não há o que sintetizar.
    """
    limite = min(LINHAS_DE_BUSCA_DO_CABECALHO, planilha.max_row)
    for linha in range(1, limite + 1):
        rotulos = {
            str(planilha.cell(linha, coluna).value or "").strip().lower()
            for coluna in range(1, planilha.max_column + 1)
        }
        if rotulos & COLUNAS_IDENTIFICADORAS:
            return linha
    return None


def _sintetizar_abas(livro: Any) -> list[str]:
    """Troca as colunas identificadoras por valores gerados, no lugar.

    Substituir no lugar, em vez de reconstruir a aba, é o que garante que a
    fixture continue com a formatação, as mesclagens e a contagem de linhas da
    original — que é exatamente o que os testes dos anexos precisam exercitar.
    """
    tratadas = []
    for aba in ABAS_COM_DADO_PESSOAL:
        if aba not in livro.sheetnames:
            continue
        planilha = livro[aba]
        linha_cabecalho = _linha_do_cabecalho(planilha)
        if linha_cabecalho is None:
            continue
        alvo = {
            coluna: str(planilha.cell(linha_cabecalho, coluna).value or "").strip().lower()
            for coluna in range(1, planilha.max_column + 1)
        }
        alvo = {c: rotulo for c, rotulo in alvo.items() if rotulo in COLUNAS_IDENTIFICADORAS}

        for sequencia, linha in enumerate(range(linha_cabecalho + 1, planilha.max_row + 1)):
            for coluna, rotulo in alvo.items():
                celula = planilha.cell(linha, coluna)
                if celula.value is None:
                    continue
                if rotulo == "nome":
                    celula.value = _nome_sintetico(sequencia)
                elif rotulo in ("email", "e-mail"):
                    celula.value = f"{_login_sintetico(sequencia)}@{DOMINIO_SINTETICO}"
                else:
                    celula.value = _login_sintetico(sequencia)

        tratadas.append(f"{aba} ({len(alvo)} coluna(s))")
    return tratadas


def sanitizar_planilha(
    nome_de_origem: str = PLANILHA, nome_de_destino: str = "levantamento.xlsx"
) -> None:
    origem = ORIGEM / nome_de_origem
    if not origem.exists():
        sys.exit(f"planilha original não encontrada: {origem}")

    # data_only=True e obrigatorio: varias celulas da aba Levantamento sao
    # formulas. Carregando sem ele, o openpyxl preserva a formula e descarta o
    # valor em cache — e como ele nao avalia formulas, a fixture sairia com
    # celulas vazias. A USN (14.070.00001.00) foi o caso que expos isso.
    repetidas = _linhas_repetidas_na_origem(origem)
    livro = openpyxl.load_workbook(origem, data_only=True)
    tratadas = _sintetizar_abas(livro)
    faixas = _materializar_faixas(livro, repetidas)

    destino = DESTINO / nome_de_destino
    livro.save(destino)
    livro.close()
    _conferir(destino)
    _conferir_dado_pessoal(destino)
    # Saida em ASCII: o console do Windows usa cp1252 e quebra com emoji/travessao.
    print(
        f"[ok] {destino.name}: abas sintetizadas - {', '.join(tratadas)}"
        f" | {faixas} mesclagem(ns) materializada(s) em {ABA_DE_ITENS}"
    )


def _conferir(destino: Path) -> None:
    """Autoconferencia: a sanitizacao nao pode esvaziar celula nenhuma.

    Sem isto, uma perda de valor passa silenciosamente para os testes e vira
    quantidade zero no relatorio.
    """
    livro = openpyxl.load_workbook(destino, read_only=True, data_only=True)
    try:
        aba = livro["Levantamento"]
        vazias = [
            linha[1].row
            for linha in aba.iter_rows(min_col=1, max_col=5)
            if linha[0].value is not None
            and any(c.value is not None for c in linha[1:4])
            and linha[4].value is None
        ]
    finally:
        livro.close()

    if vazias:
        sys.exit(
            f"[erro] {len(vazias)} linha(s) com medicao vazia apos a sanitizacao: {vazias[:8]}\n"
            "       provavel perda de valor de formula - confira o data_only na leitura"
        )


def _conferir_dado_pessoal(destino: Path) -> None:
    """Autoconferencia: nenhum endereco institucional sobrevive a sintetizacao.

    O hook de pre-commit ja bloqueia a planilha integra. Isto pega o outro
    caminho, que e o mais provavel: a fixture ser regerada com a lista de
    colunas desatualizada, e o dado real passar sem ninguem notar.
    """
    livro = openpyxl.load_workbook(destino, read_only=True, data_only=True)
    try:
        vazamentos = []
        for aba in ABAS_COM_DADO_PESSOAL:
            if aba not in livro.sheetnames:
                continue
            for linha in livro[aba].iter_rows(values_only=True):
                for valor in linha:
                    if isinstance(valor, str) and "@" in valor and DOMINIO_SINTETICO not in valor:
                        vazamentos.append(f"{aba}: {valor}")
    finally:
        livro.close()

    if vazamentos:
        sys.exit(
            f"[erro] {len(vazamentos)} valor(es) com endereco real apos a sintetizacao: "
            f"{vazamentos[:3]}\n"
            "       confira COLUNAS_IDENTIFICADORAS e a linha de cabecalho de cada aba"
        )


def copiar_pdfs() -> None:
    """T-06 — Contrato e relatório modelo são documentos públicos de contratação."""
    for origem_nome, destino_nome in (
        (CONTRATO, "contrato.pdf"),
        (MODELO, "modelo.pdf"),
        (CONTRATO_PGM, "contrato_pgm.pdf"),
        (ADITIVO_PGM, "aditivo_pgm.pdf"),
    ):
        origem = ORIGEM / origem_nome
        if not origem.exists():
            sys.exit(f"arquivo não encontrado: {origem}")
        shutil.copy2(origem, DESTINO / destino_nome)
        print(f"[ok] {destino_nome}")


def main() -> None:
    DESTINO.mkdir(parents=True, exist_ok=True)
    sanitizar_planilha()
    sanitizar_planilha(PLANILHA_PGM, "levantamento_pgm.xlsx")
    copiar_pdfs()


if __name__ == "__main__":
    main()
